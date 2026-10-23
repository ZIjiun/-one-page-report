#!/usr/bin/env python3
"""Build an editable 16:9 PowerPoint version of a one-page report spec.

Usage:
  python3 build_pptx.py SPEC.json --out report.pptx --png report.png                 # picture (default)
  python3 build_pptx.py SPEC.json --out report.pptx --html report.html --editable    # editable

Two modes. Default is picture; pass --editable to try native:
  native   python-pptx + a browser: the HTML is measured in the browser and every box,
           line of text and icon is placed at the same spot as in the PNG. Text is
           editable and charts are NATIVE (right-click > Edit Data works).
  picture  python-pptx or a browser is missing: a minimal .pptx is written with the standard library,
           holding the rendered PNG as a full-slide picture (needs --png). Not editable,
           but it always opens in PowerPoint, Keynote and Google Slides.

Chart approach adapted from the idea in enterprise-ai-skills / mckinsey-charts (MIT,
Sruthi Reddy): one highlight colour, everything else grey, no gridlines, claim titles.
"""
import argparse, json, os, re, sys, zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_report as br  # noqa: E402

SLIDE_W, SLIDE_H = 12192000, 6858000      # 13.333 x 7.5 in (16:9), EMU
FONT_FIT = 0.94                           # leave ~6% room: PowerPoint/LibreOffice set text wider than browsers
PX = SLIDE_W / 1920                       # EMU per CSS px of the HTML layout


def plain(s):
    """Strip the small markup set, keep the words."""
    s = re.sub(r"\{\{icon:[a-z0-9-]+\}\}", "", str(s or ""))
    return re.sub(r"(\*\*|==|!!|\^\^|\+\+)", "", s)


def runs(s):
    """Split markup into (text, bold, colour_role) runs."""
    s = re.sub(r"\{\{icon:[a-z0-9-]+\}\}", "", str(s or ""))
    out, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*|==(.+?)==|!!(.+?)!!|\^\^(.+?)\^\^|\+\+(.+?)\+\+", s):
        if m.start() > pos:
            out.append((s[pos:m.start()], False, None))
        if m.group(1): out.append((m.group(1), True, None))
        elif m.group(2): out.append((m.group(2), True, "good"))
        elif m.group(3): out.append((m.group(3), True, "bad"))
        elif m.group(4): out.append((m.group(4), True, "hi"))
        else: out.append((m.group(5), True, "good"))
        pos = m.end()
    if pos < len(s):
        out.append((s[pos:], False, None))
    return out


# ---------------------------------------------------------------- picture mode (stdlib only)
def picture_pptx(png, out):
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
          '<Default Extension="png" ContentType="image/png"/>'
          '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
          '<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
          '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>'
          '<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>'
          '<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/></Types>')
    R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    A = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    rel = lambda rid, typ, tgt: f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{typ}" Target="{tgt}"/>'
    rels = lambda *x: '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">' + "".join(x) + "</Relationships>"
    empty_tree = '<p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree></p:cSld>'
    theme = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><a:theme {A} name="Office"><a:themeElements><a:clrScheme name="Office">'
             + "".join(f'<a:{n}><a:srgbClr val="{v}"/></a:{n}>' for n, v in [("dk1", "000000"), ("lt1", "FFFFFF"), ("dk2", "1F497D"), ("lt2", "EEECE1"), ("accent1", "2A78D6"), ("accent2", "C0504D"), ("accent3", "9BBB59"), ("accent4", "8064A2"), ("accent5", "4BACC6"), ("accent6", "F79646"), ("hlink", "0000FF"), ("folHlink", "800080")])
             + '</a:clrScheme><a:fontScheme name="Office"><a:majorFont><a:latin typeface="Calibri"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont><a:minorFont><a:latin typeface="Calibri"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont></a:fontScheme>'
             '<a:fmtScheme name="Office"><a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst>'
             '<a:lnStyleLst><a:ln w="9525"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="25400"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="38100"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst>'
             '<a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst>'
             '<a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst></a:fmtScheme></a:themeElements></a:theme>')
    files = {
        "[Content_Types].xml": ct,
        "_rels/.rels": rels(rel("rId1", "officeDocument", "ppt/presentation.xml")),
        "ppt/presentation.xml": (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:presentation {A} {R} {P}><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>'
                                 f'<p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst><p:sldSz cx="{SLIDE_W}" cy="{SLIDE_H}"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>'),
        "ppt/_rels/presentation.xml.rels": rels(rel("rId1", "slideMaster", "slideMasters/slideMaster1.xml"), rel("rId2", "slide", "slides/slide1.xml"), rel("rId3", "theme", "theme/theme1.xml")),
        "ppt/slideMasters/slideMaster1.xml": (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:sldMaster {A} {R} {P}>{empty_tree}'
                                              '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
                                              '<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst></p:sldMaster>'),
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": rels(rel("rId1", "slideLayout", "../slideLayouts/slideLayout1.xml"), rel("rId2", "theme", "../theme/theme1.xml")),
        "ppt/slideLayouts/slideLayout1.xml": f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:sldLayout {A} {R} {P} type="blank">{empty_tree}</p:sldLayout>',
        "ppt/slideLayouts/_rels/slideLayout1.xml.rels": rels(rel("rId1", "slideMaster", "../slideMasters/slideMaster1.xml")),
        "ppt/theme/theme1.xml": theme,
        "ppt/slides/slide1.xml": (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:sld {A} {R} {P}><p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
                                  '<p:pic><p:nvPicPr><p:cNvPr id="2" name="Report"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="rId2"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
                                  f'<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{SLIDE_W}" cy="{SLIDE_H}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic></p:spTree></p:cSld></p:sld>'),
        "ppt/slides/_rels/slide1.xml.rels": rels(rel("rId1", "slideLayout", "../slideLayouts/slideLayout1.xml"), rel("rId2", "image", "../media/image1.png")),
    }
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in files.items():
            z.writestr(k, v)
        z.write(png, "ppt/media/image1.png")


# ---------------------------------------------------------------- native mode (python-pptx)
def measure(html_path):
    """Open the rendered HTML in a browser and read the real position of every piece."""
    from playwright.sync_api import sync_playwright
    import render
    js = open(os.path.join(HERE, "layout_extract.js"), encoding="utf-8").read()
    with sync_playwright() as p:
        try:
            b = p.chromium.launch()
        except Exception:
            exe = render.find_browser()
            if not exe:
                raise
            b = p.chromium.launch(executable_path=exe)
        pg = b.new_page(viewport={"width": 1920, "height": 1080})
        pg.goto("file://" + os.path.abspath(html_path))
        pg.evaluate("document.fonts.ready"); pg.wait_for_timeout(400)
        pg.evaluate(render.FIT_JS); pg.wait_for_timeout(300)
        data = pg.evaluate(js)
        # rasterise each SVG piece (icons, timeline) to PNG at 3x for crisp placement
        pngs = {}
        for k, it in enumerate(data["items"]):
            if it["kind"] in ("icon", "timeline") or (it["kind"] == "chart"):
                pg2 = b.new_page(viewport={"width": max(1, int(it["w"]) + 2), "height": max(1, int(it["h"]) + 2)}, device_scale_factor=3)
                pg2.set_content(f'<html><body style="margin:0;background:transparent">{it["svg"]}</body></html>')
                pngs[k] = pg2.screenshot(omit_background=True, clip={"x": 0, "y": 0, "width": it["w"], "height": it["h"]})
                pg2.close()
        b.close()
    return data, pngs


def native_pptx(spec, style, out, html_path):
    import io
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

    from pptx.oxml.ns import qn
    data, pngs = measure(html_path)
    lang = "zh" if spec.get("lang", "zh").startswith("zh") else "en"
    bundled = style["fonts"][lang].get("bundled")
    fam = re.findall(r'"([^"]+)"|([A-Za-z][\w ]+)', style["fonts"][lang]["family"])
    font = next((a or b for a, b in fam if (a or b).strip()), "Calibri")
    # bundled web fonts are not installed on most PCs: use a common system face in PowerPoint
    # Bundled web fonts may be missing on the viewer's PC. Use the face the HTML used when it is a
    # common install (Inter), otherwise a widely installed one with similar width.
    font = {"lxgw-wenkai-tc": "Microsoft JhengHei", "patrick-hand": "Segoe Print", "inter": "Inter"}.get(bundled, font)
    ea_font = "Microsoft JhengHei"
    rgb = lambda h: RGBColor.from_string(h.lstrip("#")[:6])
    E = lambda v: Emu(int(round(v * PX)))
    ch = style["chart"]

    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(SLIDE_W), Emu(SLIDE_H)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb(data["page"]["bg"])
    items = data["items"]

    # boxes first (zones under cards under chips), in document order which is already back-to-front
    for it in items:
        if it["kind"] != "box":
            continue
        if not it["fill"] and not it["line"]:
            continue
        shp_t = MSO_SHAPE.OVAL if it.get("circle") else (MSO_SHAPE.ROUNDED_RECTANGLE if it["round"] else MSO_SHAPE.RECTANGLE)
        shp = s.shapes.add_shape(shp_t, E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]))
        if shp_t == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = min(0.5, 10 / max(8, min(it["w"], it["h"])))
        if it["fill"]:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(it["fill"])
        else:
            shp.fill.background()
        if it["line"]:
            shp.line.color.rgb = rgb(it["line"]); shp.line.width = Pt(max(0.5, it["lw"] * 0.75))
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if it.get("accent"):
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(it["x"]), E(it["y"]), E(4), E(it["h"]))
            bar.fill.solid(); bar.fill.fore_color.rgb = rgb(it["accent"]); bar.line.fill.background(); bar.shadow.inherit = False

    # charts: native, editable
    for k, it in enumerate(items):
        if it["kind"] != "chart":
            continue
        b = json.loads(it["data"]); kind_ = b.get("kind", "bar")
        cd = CategoryChartData(); cd.categories = [d["label"] for d in b["data"]]
        cd.add_series(plain(b.get("title", "Value")), [d["value"] for d in b["data"]])
        kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}[kind_]
        gf = s.shapes.add_chart(kind, E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]), cd); c = gf.chart
        c.has_legend = False; c.has_title = False; c.font.size = Pt(8); c.font.name = font
        va = c.value_axis; va.has_major_gridlines = False; va.visible = False
        ca = c.category_axis; ca.tick_labels.font.size = Pt(8); ca.format.line.color.rgb = rgb(ch["axis"])
        if kind == XL_CHART_TYPE.BAR_CLUSTERED:
            ca.reverse_order = True
        plot = c.plots[0]; plot.has_data_labels = True
        dl = plot.data_labels; dl.font.size = Pt(8.5); dl.font.bold = True; dl.number_format = "General"; dl.number_format_is_linked = False
        ser = plot.series[0]
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            dl.position = XL_LABEL_POSITION.ABOVE
            ser.format.line.color.rgb = rgb(ch["highlight"]); ser.format.line.width = Pt(2.25); ser.smooth = False
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(ch["highlight"]); ser.marker.format.line.color.rgb = rgb("#ffffff")
            vals = [d["value"] for d in b["data"]] + ([b["reference"]["value"]] if b.get("reference") else [])
            va.minimum_scale = b.get("y_min", min(vals) - 5); va.maximum_scale = b.get("y_max", max(vals) + 3)
        else:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END; plot.gap_width = 55
            for i, d in enumerate(b["data"]):
                pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(ch["highlight"] if d.get("highlight") else ch["context"])
        # use the display text (e.g. "12.1 萬") as the data label
        for i, d in enumerate(b["data"]):
            if "display" in d:
                tf = ser.points[i].data_label.text_frame; tf.text = str(d["display"])
                r0 = tf.paragraphs[0].runs[0]; r0.font.size = Pt(8.5); r0.font.bold = True
                r0.font.color.rgb = rgb(ch["strong"] if d.get("highlight") or (kind == XL_CHART_TYPE.LINE_MARKERS and i == len(b["data"]) - 1) else ch["text"])
                if kind == XL_CHART_TYPE.LINE_MARKERS:
                    ser.points[i].data_label.position = XL_LABEL_POSITION.ABOVE
                else:
                    ser.points[i].data_label.position = XL_LABEL_POSITION.OUTSIDE_END

    # icons and timelines: pictures (vector look kept at 3x)
    chart_boxes = [it for it in items if it["kind"] == "chart"]
    def inside_chart(t):
        cx, cy = t["x"] + t["w"] / 2, t["y"] + t["h"] / 2
        return any(c["x"] <= cx <= c["x"] + c["w"] and c["y"] <= cy <= c["y"] + c["h"] for c in chart_boxes)
    for k, it in enumerate(items):
        if it["kind"] in ("icon", "timeline") and k in pngs:
            s.shapes.add_picture(io.BytesIO(pngs[k]), E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]))

    # text, one box per visual line, exactly where the browser put it
    for it in items:
        if it["kind"] != "text" or inside_chart(it):
            continue
        from pptx.enum.text import MSO_AUTO_SIZE
        # Width: up to the next element on the same line, so a slightly wider fallback font
        # can grow without colliding (shrink-on-overflow handles the rest).
        room = it["w"] * 1.08 + 4
        for o in items:
            if o is it or o["kind"] not in ("text", "chart", "icon", "box") or o["kind"] == "box":
                continue
            if o["x"] > it["x"] + it["w"] - 1 and abs((o["y"] + o["h"] / 2) - (it["y"] + it["h"] / 2)) < it["h"] / 2:
                room = min(room, o["x"] - it["x"] - 2)
        room = max(room, it["w"])
        tb = s.shapes.add_textbox(E(it["x"]), E(it["y"]), E(room), E(it["h"]))
        tf = tb.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        # Ask PowerPoint to shrink the text if a substitute font is wider than the measured one.
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        for ch_ in list(bodyPr):
            if ch_.tag in (qn("a:spAutoFit"), qn("a:noAutofit"), qn("a:normAutofit")):
                bodyPr.remove(ch_)
        bodyPr.append(bodyPr.makeelement(qn("a:normAutofit"), {}))
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Emu(0))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        for run in it["runs"]:
            r = p.add_run(); r.text = run["text"].replace("  ", " \u00a0")
            r.font.size = Pt(run["size"] * 0.75 * FONT_FIT); r.font.bold = run["bold"]; r.font.name = font
            rPr = r._r.get_or_add_rPr()
            ea = rPr.find(qn("a:ea"))
            if ea is None:
                ea = rPr.makeelement(qn("a:ea"), {}); rPr.append(ea)
            ea.set("typeface", ea_font)
            if run.get("color"):
                r.font.color.rgb = rgb(run["color"])
    prs.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec"); ap.add_argument("--out", required=True)
    ap.add_argument("--style"); ap.add_argument("--png", help="rendered PNG, used for picture mode")
    ap.add_argument("--html", help="report HTML from build_report.py, used for native mode")
    ap.add_argument("--mode", choices=["auto", "native", "picture"], default="picture",
                    help="picture (default): the PNG on one slide, never breaks. native/auto: editable text and charts")
    ap.add_argument("--editable", action="store_true", help="same as --mode auto: editable when python-pptx and a browser exist")
    a = ap.parse_args()
    spec = json.load(open(a.spec, encoding="utf-8"))
    style = br.load_style(a.style or spec.get("style", "handdrawn"))
    mode = "auto" if a.editable else a.mode
    if mode == "auto":
        try:
            import pptx  # noqa: F401
            from playwright.sync_api import sync_playwright  # noqa: F401
            mode = "native"
        except ImportError:
            mode = "picture"
    if mode == "native":
        if not a.html or not os.path.isfile(a.html):
            sys.exit("native mode needs --html (the report HTML from build_report.py)")
        native_pptx(spec, style, a.out, a.html)
    else:
        if not a.png or not os.path.isfile(a.png):
            sys.exit("picture mode needs --png (render the report with render.py first)")
        picture_pptx(a.png, a.out)
    print(json.dumps({"pptx": a.out, "mode": mode, "editable_charts": mode == "native"}))


if __name__ == "__main__":
    main()
